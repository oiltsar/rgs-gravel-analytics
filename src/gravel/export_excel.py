"""Excel с исходными данными РГС: результаты, гонщики, этапы, дистанции + сводка на формулах.

Запуск: uv run python -m gravel.export_excel  ->  exports/RGS_data_2018-2026.xlsx
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
PROC, RAW, REF = ROOT / "data" / "processed", ROOT / "data" / "raw", ROOT / "data" / "reference"
OUT = ROOT / "exports" / "RGS_data_2018-2026.xlsx"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="000000")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
THIN = Border(bottom=Side(style="thin", color="D0D0CC"))

STATUS_RU = {"finished": "финиш", "dnf": "сход", "dns": "неявка", "dsq": "дисквалификация"}
KEY_TO_EVENT = {"ЦГ": "Царь Грейдер", "ЦГ180": "Царь Грейдер", "ЦГ250": "Царь Грейдер", "FURY ROAD": "Fury Road", "FR": "Fury Road",
                "ПОКРОВА": "Покрова", "ПКРВ": "Покрова", "СПОРТМАРАФОН": "Спортмарафон Фест", "SMF": "Спортмарафон Фест",
                "MODDER": "Моддер / Ардор", "ARD": "Моддер / Ардор", "А": "Моддер / Ардор", "SHULZ GW": "SHULZ Gravel Weekend",
                "SGW": "SHULZ Gravel Weekend", "Ш": "SHULZ Gravel Weekend", "GRAVEL INSTINCT": "Gravel Instinct", "GI": "Gravel Instinct"}


def hms(sec):
    if pd.isna(sec):
        return None
    sec = int(round(sec))
    return f"{sec // 3600}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def build_frames() -> dict[str, pd.DataFrame]:
    res = pd.read_parquet(PROC / "fct_result.parquet")
    rid = pd.read_parquet(PROC / "dim_rider.parquet")
    race = pd.read_parquet(PROC / "dim_race.parquet")
    ev = pd.read_parquet(PROC / "dim_event.parquet")
    ven = pd.read_csv(REF / "venues.csv")

    in_rating = {y: {KEY_TO_EVENT[k] for e in json.loads((RAW / f"rankings_{y}.json").read_text(encoding="utf-8"))["entries"]
                     for k in e["breakdown"]} for y in (2024, 2025, 2026)}

    events = ev.merge(ven, on="event_id", how="left")
    events["date"] = pd.to_datetime(events["race_date"]).fillna(pd.to_datetime(events["start_date"]))
    events["rating"] = [("—" if y < 2024 else ("да" if n in in_rating[y] else "нет")) for y, n in zip(events.year, events.event_name)]
    events = events.sort_values(["year", "date", "event_name"])

    r = (res.merge(race[["race_id", "race_title", "distance_km", "is_long"]], on="race_id")
            .merge(events[["event_id", "event_name", "date", "macro_region"]], on="event_id")
            .merge(rid[["rider_id", "display_name"]], on="rider_id"))
    results = pd.DataFrame({
        "Год": r.year, "Этап": r.event_name, "Дата": r.date.dt.date, "Регион": r.macro_region,
        "Дистанция": r.race_title, "Км": r.distance_km, "Главная дистанция": r.is_long.map({True: "да", False: "нет"}),
        "ID гонщика": r.rider_id, "Гонщик": r.display_name, "Имя в протоколе": r.name_raw,
        "Пол": r.gender.map({"M": "М", "F": "Ж"}), "Номер": r.bib, "Команда": r.team.replace("", None), "Категория": r.category.replace("", None),
        "Статус": r.status.map(STATUS_RU), "Место в протоколе": r.place, "Очки РГС": r.points,
        "Время": r.time_s.map(hms), "Время, с": r.time_s.round(1), "Скорость, км/ч": r.speed_kmh.round(1),
        "Место среди своего пола": r.place_in_group, "Финишёров своего пола": r.finishers_in_group,
        "Перцентиль (0 = победитель)": r.pct_rank.round(3),
    }).sort_values(["Год", "Дата", "Этап", "Дистанция", "Место в протоколе"], na_position="last")

    riders = pd.DataFrame({
        "ID гонщика": rid.rider_id, "Гонщик": rid.display_name, "Пол": rid.gender.map({"M": "М", "F": "Ж"}),
        "Первый сезон": rid.first_year, "Последний сезон": rid.last_year, "Сезонов": rid.seasons,
        "Стартов": rid.starts, "Финишей": rid.finishes,
    }).sort_values(["Стартов", "Гонщик"], ascending=[False, True])

    stages = pd.DataFrame({
        "Год": events.year, "Этап": events.event_name, "Дата": events.date.dt.date, "В рейтинге РГС": events.rating,
        "Место": events.venue, "Область": events.region, "Регион": events.macro_region,
        "Широта": events.lat, "Долгота": events.lon, "Точность даты": events.date_quality, "Точность координат": events.coord_quality,
        "Источник места/даты": events.source, "Протокол": events.source_url,
    })

    dist = (race.merge(events[["event_id", "year", "event_name"]], on="event_id")
               .sort_values(["year", "event_name", "max_points"], ascending=[True, True, False]))
    distances = pd.DataFrame({
        "Год": dist.year, "Этап": dist.event_name, "Дистанция": dist.race_title, "Км": dist.distance_km,
        "Коэффициент очков": dist.max_points, "Главная дистанция": dist.is_long.map({True: "да", False: "нет"}),
        "Класс велосипеда": dist.bike_class.map({"multi": "мультиспид", "fixed": "фикс"}),
    })
    return {"Результаты": results, "Гонщики": riders, "Этапы": stages, "Дистанции": distances}


def write_table(ws, df: pd.DataFrame, widths: dict[str, int]) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for c, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, max(10, min(28, len(name) + 2)))
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    frames = build_frames()
    wb = Workbook()

    # --- лист «О файле»
    about = wb.active
    about.title = "О файле"
    lines = [
        ("Русская гравийная серия — исходные данные 2018–2026", True),
        ("Источник: открытый архив протоколов gravelseries.ru (выгрузка 04.10.2026), справочник мест и дат — сайты и каналы гонок.", False),
        ("Проект и методика: github.com/oiltsar/rgs-gravel-analytics", False),
        ("", False),
        ("Листы", True),
        ("Сводка — число стартов, финишей и сходов по годам (формулы считают по листу «Результаты»).", False),
        ("Результаты — одна строка = один старт гонщика на одной дистанции.", False),
        ("Гонщики — один человек после сведения дублей имён («ЦАРЕВ СЕРГЕЙ» = «Сергей Царев», «Дима» = «Дмитрий»).", False),
        ("Этапы — этап серии в конкретный год: дата, место, регион, входит ли в рейтинг серии.", False),
        ("Дистанции — дистанции каждого этапа: километраж, коэффициент очков, класс велосипеда.", False),
        ("", False),
        ("Как читать", True),
        ("ID гонщика связывает «Результаты» и «Гонщики». «Имя в протоколе» — как записано организатором, «Гонщик» — нормализованное имя.", False),
        ("Статус: финиш, сход, неявка, дисквалификация. Сходы и неявки фиксируются не во всех протоколах.", False),
        ("Место в протоколе — как у организатора (иногда абсолютное, иногда по полу). «Место среди своего пола» пересчитано единообразно по времени.", False),
        ("Перцентиль: 0 — победитель среди своего пола на дистанции, 1 — последний финишёр.", False),
        ("Км и скорость — только там, где дистанция указана в названии; у дистанций без километража пусто.", False),
        ("Главная дистанция — с максимальным коэффициентом очков этапа (1000).", False),
        ("Места и даты этапов до 2026 года собраны вручную; у части координат точность «approx», ссылка на источник — в листе «Этапы».", False),
        ("Тёзки (~0,3% гонщиков) могут быть сведены в одного человека; сменившие фамилию — наоборот, разделены.", False),
    ]
    for text, bold in lines:
        about.append([text])
        about.cell(row=about.max_row, column=1).font = Font(name=FONT, bold=bold, size=14 if about.max_row == 1 else 10)
    about.column_dimensions["A"].width = 130

    # --- данные
    widths = {"Этап": 22, "Дистанция": 26, "Гонщик": 28, "Имя в протоколе": 28, "Команда": 22, "Место": 34,
              "Источник места/даты": 40, "Протокол": 40, "Дата": 12}
    for name in ["Результаты", "Гонщики", "Этапы", "Дистанции"]:
        write_table(wb.create_sheet(name), frames[name], widths)

    # --- сводка на формулах (считается по листу «Результаты»)
    s = wb.create_sheet("Сводка", 1)
    s.append(["Год", "Стартов (без неявок)", "Финишей", "Сходов", "Неявок", "Доля сходов", "Этапов"])
    years = sorted(frames["Результаты"]["Год"].unique())
    n = len(frames["Результаты"]) + 1
    R = f"Результаты!$A$2:$A${n}"; ST = f"Результаты!$O$2:$O${n}"
    for i, y in enumerate(years, start=2):
        s.append([int(y),
                  f'=COUNTIFS({R},A{i},{ST},"<>неявка")',
                  f'=COUNTIFS({R},A{i},{ST},"финиш")',
                  f'=COUNTIFS({R},A{i},{ST},"сход")',
                  f'=COUNTIFS({R},A{i},{ST},"неявка")',
                  f"=IF(B{i}=0,0,D{i}/B{i})",
                  f"=COUNTIF(Этапы!$A$2:$A${len(frames['Этапы']) + 1},A{i})"])
    t = len(years) + 2
    s.append(["Итого", f"=SUM(B2:B{t-1})", f"=SUM(C2:C{t-1})", f"=SUM(D2:D{t-1})", f"=SUM(E2:E{t-1})", f"=IF(B{t}=0,0,D{t}/B{t})", f"=SUM(G2:G{t-1})"])
    for c in range(1, 8):
        h = s.cell(row=1, column=c); h.font, h.fill = HEAD_FONT, HEAD_FILL
        h.alignment = Alignment(wrap_text=True, vertical="center")
        s.column_dimensions[get_column_letter(c)].width = 16
    for row in s.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.row == t)
            if cell.column == 6:
                cell.number_format = "0.0%"
            elif cell.column > 1:
                cell.number_format = "#,##0"
    s.row_dimensions[1].height = 30
    s.append([])
    s.append(["Сходы и неявки фиксируются не во всех протоколах: в 2022–2024 годах их нет совсем, поэтому доля сходов за эти годы — ноль условно."])
    s.cell(row=s.max_row, column=1).font = Font(name=FONT, size=9, italic=True, color="74746F")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT, {k: len(v) for k, v in frames.items()})


if __name__ == "__main__":
    main()
